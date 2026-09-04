import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def highlight_min_values(file_path):
    # 加载工作簿
    wb = load_workbook(file_path)
    ws = wb.active

    # 定义黄色填充
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    # 获取列名对应的列号
    columns = {}
    for cell in ws[1]:
        columns[cell.value] = cell.column

    # 需要的列名
    delay_columns = ["GAA delay", "QC-HHO delay", "GACSM delay", "QHA delay"]
    energy_columns = ["GAA energy", "QC-HHO energy", "GACSM energy", "QHA energy"]

    # 检查所有需要的列是否存在
    for col in delay_columns + energy_columns:
        if col not in columns:
            print(f"警告: 列 '{col}' 不存在于表格中")
            return

    # 初始化统计计数器
    best_counts = {
        "GAA": 0,
        "QC-HHO": 0,
        "GACSM": 0,
        "QHA": 0
    }

    # 从第二行开始处理每一行
    for row in range(2, ws.max_row + 1):
        # 获取delay值
        delay_values = {
            col: ws.cell(row=row, column=columns[col]).value
            for col in delay_columns
        }

        # 找到最小delay值
        min_delay = min(delay_values.values())
        min_delay_cols = [col for col, val in delay_values.items() if val == min_delay]

        # 如果只有一个最小delay，直接标记
        if len(min_delay_cols) == 1:
            best_col = min_delay_cols[0]
            # 获取对应的energy列名
            energy_col = best_col.replace(" delay", " energy")

            # 标记delay和energy单元格
            ws.cell(row=row, column=columns[best_col]).fill = yellow_fill
            ws.cell(row=row, column=columns[energy_col]).fill = yellow_fill

            # 更新统计计数器
            algorithm = best_col.split()[0]  # 提取算法名称
            best_counts[algorithm] += 1
        else:
            # 如果有多个最小delay，比较energy值
            energy_values = {
                col: ws.cell(row=row, column=columns[col.replace(" delay", " energy")]).value
                for col in min_delay_cols
            }

            min_energy = min(energy_values.values())
            min_energy_cols = [col for col, val in energy_values.items() if val == min_energy]

            # 标记所有最小energy对应的delay和energy单元格
            for col in min_energy_cols:
                energy_col = col.replace(" delay", " energy")
                ws.cell(row=row, column=columns[col]).fill = yellow_fill
                ws.cell(row=row, column=columns[energy_col]).fill = yellow_fill

                # 更新统计计数器
                algorithm = col.split()[0]  # 提取算法名称
                best_counts[algorithm] += 1

    # 保存修改
    wb.save(file_path)
    print(f"处理完成，结果已保存到 {file_path}")

    # 打印统计结果
    print("\n最佳结果统计:")
    for algorithm, count in best_counts.items():
        print(f"{algorithm}: {count}次")
    print(f"总计: {sum(best_counts.values())}次")


# 使用示例
save_dict = ".\\"  # 替换为你的目录路径
excel_file = os.path.join(save_dict, "total.xlsx")

if os.path.exists(excel_file):
    highlight_min_values(excel_file)
else:
    print(f"错误: 文件 {excel_file} 不存在")