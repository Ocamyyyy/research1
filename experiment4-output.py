import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl import Workbook


def highlight_and_statistics(file_path):
    # 加载工作簿
    wb = load_workbook(file_path)
    ws = wb.active

    # 定义黄色填充
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    # 获取列名对应的列号
    columns = {}
    for cell in ws[1]:
        columns[cell.value] = cell.column

    # 需要的列名
    delay_columns = ["GAA delay", "QC-HHO delay", "GACSM delay", "QHA delay"]
    energy_columns = ["GAA energy", "QC-HHO energy", "GACSM energy", "QHA energy"]

    # 检查所有需要的列是否存在
    for col in delay_columns + energy_columns:
        if col not in columns:
            print(f"警告: 列 '{col}' 不存在于表格中")
            return

    # 初始化统计数据结构
    parameter_groups = [
        (2, 5, 30), (2, 5, 40), (2, 5, 50), (2, 5, 60), (2, 5, 70), (2, 5, 80), (2, 5, 90), (2, 5, 100),
        (2, 10, 40), (2, 10, 70), (2, 10, 100), (2, 10, 200), (2, 10, 300), (2, 10, 400), (2, 10, 500),
        (2, 20, 100), (2, 20, 200), (2, 20, 300), (2, 20, 400), (2, 20, 500)
    ]

    # 为每组参数创建一个统计字典
    group_stats = {}
    for group in parameter_groups:
        group_stats[group] = {
            "GAA": 0,
            "QC-HHO": 0,
            "GACSM": 0,
            "QHA": 0,
            "total": 0
        }

    # 从第二行开始处理每一行
    for row in range(2, ws.max_row + 1):
        # 获取N, V, tot值
        try:
            N = ws.cell(row=row, column=columns["N"]).value
            V = ws.cell(row=row, column=columns["V"]).value
            tot = ws.cell(row=row, column=columns["tot"]).value
            current_group = (N, V, tot)

            # 如果这组参数不在我们的统计范围内，跳过
            if current_group not in group_stats:
                continue

            group_stats[current_group]["total"] += 1
        except KeyError:
            print(f"警告: 行 {row} 缺少N/V/tot列")
            continue

        # 获取delay值
        delay_values = {
            col: ws.cell(row=row, column=columns[col]).value
            for col in delay_columns
        }

        # 找到最小delay值
        min_delay = min(delay_values.values())
        min_delay_cols = [col for col, val in delay_values.items() if val == min_delay]

        # 如果只有一个最小delay，直接标记
        if len(min_delay_cols) == 1:
            best_col = min_delay_cols[0]
            # 获取对应的energy列名
            energy_col = best_col.replace(" delay", " energy")

            # 标记delay和energy单元格
            ws.cell(row=row, column=columns[best_col]).fill = yellow_fill
            ws.cell(row=row, column=columns[energy_col]).fill = yellow_fill

            # 更新统计计数器
            algorithm = best_col.split()[0]  # 提取算法名称
            group_stats[current_group][algorithm] += 1
        else:
            # 如果有多个最小delay，比较energy值
            energy_values = {
                col: ws.cell(row=row, column=columns[col.replace(" delay", " energy")]).value
                for col in min_delay_cols
            }

            min_energy = min(energy_values.values())
            min_energy_cols = [col for col, val in energy_values.items() if val == min_energy]

            # 标记所有最小energy对应的delay和energy单元格
            for col in min_energy_cols:
                energy_col = col.replace(" delay", " energy")
                ws.cell(row=row, column=columns[col]).fill = yellow_fill
                ws.cell(row=row, column=columns[energy_col]).fill = yellow_fill

                # 更新统计计数器
                algorithm = col.split()[0]  # 提取算法名称
                group_stats[current_group][algorithm] += 1

    # 保存修改后的原始文件
    wb.save(file_path)
    print(f"高亮处理完成，结果已保存到 {file_path}")

    # 创建统计结果工作簿
    stat_wb = Workbook()
    stat_ws = stat_wb.active
    stat_ws.title = "算法最优解统计"

    # 写入表头
    headers = ["N", "V", "tot", "GAA(%)", "QC-HHO(%)", "GACSM(%)", "QHA(%)", "样本数"]
    stat_ws.append(headers)

    # 写入统计数据
    for group in parameter_groups:
        stats = group_stats[group]
        total = stats["total"]
        if total == 0:
            continue

        row_data = [
            group[0],  # N
            group[1],  # V
            group[2],  # tot
            round(stats["GAA"] / total * 100, 2),
            round(stats["QC-HHO"] / total * 100, 2),
            round(stats["GACSM"] / total * 100, 2),
            round(stats["QHA"] / total * 100, 2),
            total
        ]
        stat_ws.append(row_data)

    # 保存统计结果
    stat_file = os.path.join(os.path.dirname(file_path), "algorithm_statistics.xlsx")
    stat_wb.save(stat_file)
    print(f"\n统计结果已保存到 {stat_file}")

    # 打印总统计
    print("\n总体最优解统计:")
    total_counts = {
        "GAA": 0,
        "QC-HHO": 0,
        "GACSM": 0,
        "QHA": 0,
        "total": 0
    }

    for group in group_stats.values():
        for algo in ["GAA", "QC-HHO", "GACSM", "QHA"]:
            total_counts[algo] += group[algo]
        total_counts["total"] += group["total"]

    for algo in ["GAA", "QC-HHO", "GACSM", "QHA"]:
        percentage = round(total_counts[algo] / total_counts["total"] * 100, 2)
        print(f"{algo}: {total_counts[algo]}次 ({percentage}%)")
    print(f"总计: {total_counts['total']}次")


# 使用示例
save_dict = ".\\"  # 替换为你的目录路径
excel_file = os.path.join(save_dict, "total.xlsx")

if os.path.exists(excel_file):
    highlight_and_statistics(excel_file)
else:
    print(f"错误: 文件 {excel_file} 不存在")