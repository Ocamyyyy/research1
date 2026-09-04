import json
import math
import random
import sys
import time
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from GA import genetic_algorithm
from HHO import harris_hawks_optimization
from QHA import quantum_inspired_hyper_heuristics_algorithm
from SAFLA_adapted import safla_adapted
from eFLS_adapted import AdaptedEFLSMappingSolution
from solution_utils import MappingSolution, calc_subtask_instance_from_position
from task_definition import TaskConfig


TEST_LIST = [
    (2, 10, 40),
    (2, 10, 70),
    (2, 10, 100),
    (2, 10, 200),
    (2, 10, 300),
    (2, 10, 400),
    (2, 10, 500),
]
REPEAT_COUNT = 5
RESULT_ROOT = Path("results") / "experiments" / "try3"
CASE_ROOT = RESULT_ROOT / "cases"

ALGORITHM_ORDER = [
    "HEFT",
    "SAFLA_adapted",
    "eFLS_adapted",
    "GA",
    "HHO",
    "QHA",
]

PLOT_ALGORITHMS = [
    ("HEFT", "HEFT"),
    ("SAFLA_adapted", "SAFLA"),
    ("eFLS_adapted", "eFLS"),
    ("GA", "GAA"),
    ("HHO", "QC-HHO"),
    ("QHA", "AEHS"),
]

HIGHLIGHT_ALGORITHMS = [
    "SAFLA_adapted",
    "eFLS_adapted",
    "GA",
    "HHO",
    "QHA",
]

PLOT_COLORS = [
    "#F56E68",
    "#FEB5AF",
    "#FFF4F2",
    "#1597A5",
    "#0E606C",
    "#7FAE55",
]

PLOT_HATCHES = ["///", "\\\\\\", "++", "xx", "..", ""]


def build_random_dict(N, V, tot):
    return {
        "N": N,
        "V": V,
        "alpha": 1.0,
        "max_out": 2,
        "beta": 1,
        "Tpr": 100,
        "Tpr_beta": 50,
        "wDAG": 400,
        "gamma_GPU": 2.0,
        "gamma_hw": 1.5,
        "npu_eligibility_prob": 0.5,
        "fpga_eligibility_prob": 0.5,
        "P_sw": 0.5,
        "P_npu": 5.0,
        "P_npu_beta": 1.0,
        "P_static_min": 0.5,
        "P_static_max": 4.0,
        "P_dyn_min": 1.0,
        "P_dyn_max": 6.0,
        "NWR": 3.0,
        "tot": tot,
    }


def stable_seed(N, V, tot, iterr):
    return 200000 + N * 10000 + V * 100 + tot + iterr * 1000000


def case_name(N, V, tot, iterr):
    return f"N{N}_V{V}_tot{tot}_iter{iterr}"


def atomic_write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    tmp_path.replace(path)


def load_case_payload(path, N, V, tot, iterr, seed):
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {
        "case": {
            "N": N,
            "V": V,
            "tot": tot,
            "iterr": iterr,
            "seed": seed,
        },
        "algorithms": {},
        "status": "running",
    }


def is_case_fully_completed(payload):
    if payload.get("status") != "completed":
        return False
    algorithms = payload.get("algorithms", {})
    return all(name in algorithms for name in ALGORITHM_ORDER)


def build_task(N, V, tot, iterr):
    seed = stable_seed(N, V, tot, iterr)
    random.seed(seed)
    np.random.seed(seed)
    task = TaskConfig()
    task.init_class_vars(build_random_dict(N, V, tot), mode="random")
    return task, seed


def serialize_taskset(task):
    applications = []
    expanded_tasks = []

    for app_index, app in enumerate(task.app_list):
        depth, level_nodes = app.bfs_depth_and_levels()
        app_summary = {
            "app_index": app_index,
            "node_num": app.node_num,
            "period": app.period,
            "depth": depth,
            "level_nodes": {str(level): nodes for level, nodes in level_nodes.items()},
            "edges": [list(edge) for edge in app.edges],
            "sources": [node for node, degree in enumerate(app.into_degree) if degree == 0],
            "sinks": app.sink_nodes(),
            "cpu_eligible": list(app.cpu_eligible),
            "npu_eligible": list(app.npu_eligible),
            "hw_eligible": list(app.hw_eligible),
            "sw_exe": list(app.sw_exe),
            "npu_exe": list(app.npu_exe),
            "hw_exe": list(app.hw_exe),
            "tpr": list(app.tpr),
            "P_sw": list(app.P_sw),
            "P_npu": list(app.P_npu),
            "P_static": list(app.P_static),
            "P_dyn": list(app.P_dyn),
            "shape": {
                "layer_count": depth,
                "width_per_level": {
                    str(level): len(nodes) for level, nodes in level_nodes.items()
                },
            },
        }
        applications.append(app_summary)

        for instance_index in range(1, task.instance_number_list[app_index] + 1):
            arr = task.instance_dict_list[app_index][instance_index]["arr"]
            ddl = task.instance_dict_list[app_index][instance_index]["ddl"]
            for node_index in range(app.node_num):
                task_id = sum(
                    task.instance_number_list[prev] * task.app_list[prev].node_num
                    for prev in range(app_index)
                ) + (instance_index - 1) * app.node_num + node_index
                expanded_tasks.append(
                    {
                        "task_id": task_id,
                        "app_index": app_index,
                        "instance_index": instance_index,
                        "node_index": node_index,
                        "arrival_time": arr,
                        "deadline": ddl,
                        "cpu_eligible": bool(app.cpu_eligible[node_index]),
                        "npu_eligible": bool(app.npu_eligible[node_index]),
                        "hw_eligible": bool(app.hw_eligible[node_index]),
                        "predecessors": app.predecessor_of_node(node_index),
                        "successors": app.successor_of_node(node_index),
                        "sw_exe": app.sw_exe[node_index],
                        "npu_exe": app.npu_exe[node_index],
                        "hw_exe": app.hw_exe[node_index],
                        "reconfiguration_time": app.tpr[node_index],
                        "P_sw": app.P_sw[node_index],
                        "P_npu": app.P_npu[node_index],
                        "P_static": app.P_static[node_index],
                        "P_dyn": app.P_dyn[node_index],
                    }
                )

    return {
        "H": task.H,
        "E_max": task.E_max,
        "target_NWR": task.target_NWR,
        "actual_NWR": task.actual_NWR,
        "periods": list(task.periods),
        "processor_config": {
            "CPU_number": task.CPU_number,
            "NPU_number": task.NPU_number,
            "RP_number": task.RP_number,
            "processor_number": task.processor_number,
        },
        "instance_number_list": list(task.instance_number_list),
        "task_number": task.task_number,
        "applications": applications,
        "expanded_tasks": expanded_tasks,
    }


def save_task_graph_images(task, graph_dir):
    graph_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for app_index, app in enumerate(task.app_list):
        fig_path = graph_dir / f"app{app_index}_dag_shape.png"
        graph = nx.DiGraph()
        graph.add_nodes_from(range(app.node_num))
        graph.add_edges_from(app.edges)
        plt.figure(figsize=(8, 6))
        nx.draw_networkx(
            graph,
            arrows=True,
            pos=app.position,
            node_color="#DFE1E2",
            edgecolors="black",
            node_size=900,
            font_size=10,
        )
        plt.axis("off")
        plt.tight_layout()
        plt.savefig(fig_path, dpi=300)
        plt.close()
        paths.append(str(fig_path))
    return paths


def serialize_solution_schedule(task, solution):
    schedule_rows = []
    for task_id in range(task.task_number):
        app_index, instance_index, node_index = calc_subtask_instance_from_position(task, task_id)
        processor = solution.task_mapping[task_id]
        processor_type = "CPU"
        if processor >= task.CPU_number + task.NPU_number:
            processor_type = "RP"
        elif processor >= task.CPU_number:
            processor_type = "NPU"
        schedule_rows.append(
            {
                "task_id": task_id,
                "app_index": app_index,
                "instance_index": instance_index,
                "node_index": node_index,
                "processor": processor,
                "processor_type": processor_type,
                "start_time": solution.task_start_time[task_id],
                "end_time": solution.task_end_time[task_id],
                "execution_time": solution.task_end_time[task_id] - solution.task_start_time[task_id],
                "reconfigured": bool(solution.task_DPR[task_id]),
                "reconfiguration_start_time": solution.task_DPR_start_time[task_id],
                "reconfiguration_end_time": solution.task_DPR_end_time[task_id],
                "reconfiguration_time": solution.task_DPR_end_time[task_id] - solution.task_DPR_start_time[task_id]
                if solution.task_DPR[task_id]
                else 0,
            }
        )

    return {
        "processor_completed_task": [list(items) for items in solution.processor_completed_task],
        "processor_RP_time_list": [list(item) for item in solution.processor_RP_time_list],
        "tasks": schedule_rows,
    }


def save_schedule_plot(task, solution, algorithm_name, artifact_dir):
    artifact_dir.mkdir(parents=True, exist_ok=True)
    save_base = artifact_dir / f"{algorithm_name}_schedule"
    solution.visualize_init(task, window_name=str(save_base), is_block=False, is_save=True)
    return str(save_base.with_suffix(".png"))


def seed_algorithm(case_seed, algorithm_name):
    offsets = {
        "HEFT": 11,
        "SAFLA_adapted": 23,
        "eFLS_adapted": 37,
        "GA": 51,
        "HHO": 67,
        "QHA": 83,
    }
    seed = case_seed + offsets[algorithm_name]
    random.seed(seed)
    np.random.seed(seed)
    return seed


def collect_solution_metrics(task, solution, runtime_seconds, feasible=None, schedule_plot_path=None):
    energy = solution.calc_energy(task)
    delay = solution.calc_delay(task)
    if feasible is None:
        feasible = delay == 0
    return {
        "energy": float(energy),
        "energy_percent": float(energy / task.E_max * 100),
        "delay": float(delay),
        "feasible": bool(feasible),
        "runtime_seconds": float(runtime_seconds),
        "schedule_plot": schedule_plot_path,
        "schedule": serialize_solution_schedule(task, solution),
    }


def run_heft(task, case_seed, artifact_dir):
    seed_algorithm(case_seed, "HEFT")
    start = time.time()
    solution = MappingSolution(
        task.CPU_number,
        task.NPU_number,
        task.RP_number,
        task.task_number,
    )
    solution.heft_scheduling(task)
    schedule_plot_path = save_schedule_plot(task, solution, "HEFT", artifact_dir)
    return collect_solution_metrics(task, solution, time.time() - start, schedule_plot_path=schedule_plot_path)


def run_safla(task, case_seed, artifact_dir):
    seed_algorithm(case_seed, "SAFLA_adapted")
    start = time.time()
    solution, _history = safla_adapted(task, verbose=False)
    schedule_plot_path = save_schedule_plot(task, solution, "SAFLA_adapted", artifact_dir)
    return collect_solution_metrics(task, solution, time.time() - start, schedule_plot_path=schedule_plot_path)


def run_efls(task, case_seed, artifact_dir):
    seed_algorithm(case_seed, "eFLS_adapted")
    start = time.time()
    solution = AdaptedEFLSMappingSolution(
        task.CPU_number,
        task.NPU_number,
        task.RP_number,
        task.task_number,
        verbose=False,
    )
    feasible = solution.efls_scheduling(task)
    schedule_plot_path = save_schedule_plot(task, solution, "eFLS_adapted", artifact_dir)
    return collect_solution_metrics(
        task,
        solution,
        time.time() - start,
        feasible=feasible,
        schedule_plot_path=schedule_plot_path,
    )


def run_ga(task, artifact_dir, msg, iterr, case_seed):
    seed_algorithm(case_seed, "GA")
    start = time.time()
    solution, _best_energy = genetic_algorithm(
        task,
        save_dic=str(artifact_dir),
        msg=msg,
        iterr=iterr,
        ecs_initial=False,
        seed=24 + iterr,
        pop_size=50,
        max_iter=1000,
        crossover_rate=0.8,
        mutation_rate=0.1,
        plot_number=3,
    )
    return collect_solution_metrics(task, solution, time.time() - start)


def run_hho(task, artifact_dir, msg, iterr, case_seed):
    seed_algorithm(case_seed, "HHO")
    start = time.time()
    solution, _best_energy = harris_hawks_optimization(
        task,
        n=50,
        save_dic=str(artifact_dir),
        msg=msg,
        iterr=iterr,
        max_iter=1000,
        ecs_intial=False,
        is_modified=True,
    )
    return collect_solution_metrics(task, solution, time.time() - start)


def run_qha(task, artifact_dir, msg, iterr, case_seed):
    seed_algorithm(case_seed, "QHA")
    start = time.time()
    solution, _best_energy = quantum_inspired_hyper_heuristics_algorithm(
        task,
        pop_size=50,
        save_dic=str(artifact_dir),
        msg=msg,
        iterr=iterr,
        max_iter=1000,
        phase_iter=10,
        ecs_initial=False,
    )
    return collect_solution_metrics(task, solution, time.time() - start)


def ensure_case_metadata(task, payload, case_dir):
    artifacts = payload.setdefault("artifacts", {})
    if "taskset" not in payload:
        payload["taskset"] = serialize_taskset(task)
    if "task_graph_images" not in artifacts:
        artifacts["task_graph_images"] = save_task_graph_images(task, case_dir / "task_graphs")
    return payload


def ensure_lightweight_schedule_artifacts(task, payload, artifact_root, seed):
    rerun_targets = []
    for algorithm_name in ["HEFT", "SAFLA_adapted", "eFLS_adapted"]:
        info = payload.get("algorithms", {}).get(algorithm_name)
        if not info:
            continue
        if info.get("schedule_plot") and info.get("schedule"):
            continue
        rerun_targets.append(algorithm_name)

    for algorithm_name in rerun_targets:
        print(f"  [BACKFILL] {algorithm_name} schedule details/plot")
        if algorithm_name == "HEFT":
            result = run_heft(task, seed, artifact_root / "HEFT")
        elif algorithm_name == "SAFLA_adapted":
            result = run_safla(task, seed, artifact_root / "SAFLA_adapted")
        else:
            result = run_efls(task, seed, artifact_root / "eFLS_adapted")

        original = payload["algorithms"][algorithm_name]
        original["schedule_plot"] = result.get("schedule_plot")
        original["schedule"] = result.get("schedule")
        if "feasible" not in original:
            original["feasible"] = result.get("feasible")

    return payload


def run_case(N, V, tot, iterr):
    task, seed = build_task(N, V, tot, iterr)
    name = case_name(N, V, tot, iterr)
    case_dir = CASE_ROOT / name
    artifact_root = case_dir / "artifacts"
    payload_path = case_dir / "result.json"
    payload = load_case_payload(payload_path, N, V, tot, iterr, seed)
    payload = ensure_case_metadata(task, payload, case_dir)
    payload = ensure_lightweight_schedule_artifacts(task, payload, artifact_root, seed)
    atomic_write_json(payload_path, payload)

    if is_case_fully_completed(payload):
        print(f"[CASE] {name} already completed, regenerate summary/figure only")
        return payload

    print(
        f"[CASE] {name} | H={task.H} | E_max={task.E_max:.2f} | "
        f"completed={list(payload['algorithms'].keys())}"
    )

    runners = {
        "HEFT": lambda: run_heft(task, seed, artifact_root / "HEFT"),
        "SAFLA_adapted": lambda: run_safla(task, seed, artifact_root / "SAFLA_adapted"),
        "eFLS_adapted": lambda: run_efls(task, seed, artifact_root / "eFLS_adapted"),
        "GA": lambda: run_ga(task, artifact_root / "GA", f"{N}_{V}_{tot}", iterr, seed),
        "HHO": lambda: run_hho(task, artifact_root / "HHO", f"{N}_{V}_{tot}", iterr, seed),
        "QHA": lambda: run_qha(task, artifact_root / "QHA", f"{N}_{V}_{tot}", iterr, seed),
    }

    for algorithm_name in ALGORITHM_ORDER:
        if algorithm_name in payload["algorithms"]:
            print(f"  [SKIP] {algorithm_name} already completed")
            continue

        (artifact_root / algorithm_name).mkdir(parents=True, exist_ok=True)
        print(f"  [RUN ] {algorithm_name}")
        try:
            result = runners[algorithm_name]()
            payload["algorithms"][algorithm_name] = result
            atomic_write_json(payload_path, payload)
            print(
                f"  [DONE] {algorithm_name} | "
                f"energy={result['energy_percent']:.2f}% | "
                f"delay={result['delay']:.2f} | "
                f"time={result['runtime_seconds']:.2f}s"
            )
        except Exception as exc:
            payload["status"] = "failed"
            payload["failed_algorithm"] = algorithm_name
            payload["error"] = repr(exc)
            atomic_write_json(payload_path, payload)
            raise

    payload["status"] = "completed"
    atomic_write_json(payload_path, payload)
    return payload


def collect_rows():
    rows = []
    for N, V, tot in TEST_LIST:
        for iterr in range(REPEAT_COUNT):
            payload_path = CASE_ROOT / case_name(N, V, tot, iterr) / "result.json"
            if not payload_path.exists():
                continue

            payload = json.loads(payload_path.read_text(encoding="utf-8"))
            algorithms = payload.get("algorithms", {})
            row = {
                "N": N,
                "V": V,
                "tot": tot,
                "iterr": iterr,
                "status": payload.get("status", "unknown"),
            }
            for algorithm_name in ALGORITHM_ORDER:
                metrics = algorithms.get(algorithm_name, {})
                prefix = algorithm_name
                row[f"{prefix} energy"] = metrics.get("energy_percent")
                row[f"{prefix} delay"] = metrics.get("delay")
                row[f"{prefix} feasible"] = metrics.get("feasible")
                row[f"{prefix} runtime_seconds"] = metrics.get("runtime_seconds")
            rows.append(row)
    return rows


def write_summary_files():
    rows = collect_rows()
    if not rows:
        return

    df = pd.DataFrame(rows)
    df = df.sort_values(by=["tot", "iterr"], kind="stable")
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)
    raw_csv_path = RESULT_ROOT / "experiment1_raw_results.csv"
    raw_xlsx_path = RESULT_ROOT / "experiment1_raw_results.xlsx"
    df.to_csv(raw_csv_path, index=False, encoding="utf-8-sig")
    try:
        df.to_excel(raw_xlsx_path, index=False)
    except Exception as exc:
        print(f"[WARN] Failed to write raw Excel file: {exc}")

    summary_rows = []
    for (N, V, tot), group in df.groupby(["N", "V", "tot"], sort=True):
        row = {"N": N, "V": V, "tot": tot}
        for algorithm_name in ALGORITHM_ORDER:
            energy_col = f"{algorithm_name} energy"
            delay_col = f"{algorithm_name} delay"
            runtime_col = f"{algorithm_name} runtime_seconds"
            row[f"{algorithm_name} avg_energy"] = group[energy_col].dropna().mean()
            row[f"{algorithm_name} avg_delay"] = group[delay_col].dropna().mean()
            row[f"{algorithm_name} avg_runtime_seconds"] = group[runtime_col].dropna().mean()
        summary_rows.append(row)

    summary_df = pd.DataFrame(summary_rows)
    summary_df = summary_df.sort_values(by=["tot"], kind="stable")
    summary_csv_path = RESULT_ROOT / "experiment1_summary.csv"
    summary_xlsx_path = RESULT_ROOT / "experiment1_summary.xlsx"
    summary_df.to_csv(summary_csv_path, index=False, encoding="utf-8-sig")
    try:
        summary_df.to_excel(summary_xlsx_path, index=False)
    except Exception as exc:
        print(f"[WARN] Failed to write summary Excel file: {exc}")

    try:
        highlight_best_rows(raw_xlsx_path)
    except Exception as exc:
        print(f"[WARN] Failed to highlight raw Excel file: {exc}")

    try:
        draw_result_figure(summary_df)
    except Exception as exc:
        print(f"[WARN] Failed to draw result figure: {exc}")


def highlight_best_rows(excel_path):
    if not excel_path.exists():
        return

    workbook = load_workbook(excel_path)
    worksheet = workbook.active
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    columns = {cell.value: cell.column for cell in worksheet[1]}
    delay_columns = [f"{name} delay" for name in HIGHLIGHT_ALGORITHMS]
    energy_columns = [f"{name} energy" for name in HIGHLIGHT_ALGORITHMS]

    for col_name in delay_columns + energy_columns:
        if col_name not in columns:
            raise ValueError(f"Missing column for highlight: {col_name}")

    for row in range(2, worksheet.max_row + 1):
        delay_values = {
            col_name: worksheet.cell(row=row, column=columns[col_name]).value
            for col_name in delay_columns
        }
        valid_delay_values = {
            col_name: value for col_name, value in delay_values.items() if value is not None
        }
        if not valid_delay_values:
            continue

        min_delay = min(valid_delay_values.values())
        min_delay_cols = [
            col_name for col_name, value in valid_delay_values.items() if value == min_delay
        ]

        if len(min_delay_cols) == 1:
            best_delay_col = min_delay_cols[0]
            best_energy_col = best_delay_col.replace(" delay", " energy")
            worksheet.cell(row=row, column=columns[best_delay_col]).fill = yellow_fill
            worksheet.cell(row=row, column=columns[best_energy_col]).fill = yellow_fill
            continue

        energy_values = {
            col_name: worksheet.cell(
                row=row,
                column=columns[col_name.replace(" delay", " energy")],
            ).value
            for col_name in min_delay_cols
        }
        valid_energy_values = {
            col_name: value for col_name, value in energy_values.items() if value is not None
        }
        if not valid_energy_values:
            continue

        min_energy = min(valid_energy_values.values())
        for delay_col, energy_value in valid_energy_values.items():
            if energy_value == min_energy:
                energy_col = delay_col.replace(" delay", " energy")
                worksheet.cell(row=row, column=columns[delay_col]).fill = yellow_fill
                worksheet.cell(row=row, column=columns[energy_col]).fill = yellow_fill

    highlighted_path = excel_path.with_name(f"{excel_path.stem}_highlighted.xlsx")
    workbook.save(highlighted_path)


def draw_result_figure(summary_df):
    if summary_df.empty:
        return

    plot_df = summary_df.copy()
    plot_df = plot_df.sort_values(by=["tot"], kind="stable")
    x_values = [str(int(tot)) for tot in plot_df["tot"].tolist()]

    plt.rcParams["font.family"] = "Arial"
    plt.figure(figsize=(16, 10))

    energy_series = []
    for algorithm_name, _label in PLOT_ALGORITHMS:
        energy_series.append(plot_df[f"{algorithm_name} avg_energy"].tolist())

    bar_width = 0.12
    intra_group_gap = 0.04
    group_gap = 0.35
    total_group_width = len(PLOT_ALGORITHMS) * bar_width + (len(PLOT_ALGORITHMS) - 1) * intra_group_gap
    index = np.arange(len(x_values)) * (total_group_width + group_gap)

    for i, (_algorithm_name, label) in enumerate(PLOT_ALGORITHMS):
        bar_positions = index + i * (bar_width + intra_group_gap)
        plt.bar(
            bar_positions,
            energy_series[i],
            width=bar_width,
            label=label,
            color=PLOT_COLORS[i],
            edgecolor="black",
            alpha=0.85,
            hatch=PLOT_HATCHES[i],
        )

    plt.xlabel("Task Instances in Hyperperiod", fontsize=28)
    plt.ylabel("Normalized Energy-consumption (%)", fontsize=28)
    group_center_offset = (len(PLOT_ALGORITHMS) - 1) * (bar_width + intra_group_gap) / 2
    plt.xticks(index + group_center_offset, x_values, fontsize=24)
    plt.yticks(fontsize=24)
    max_value = max(max(series) for series in energy_series if series)
    plt.ylim(0, max_value * 1.15 if max_value > 0 else 1)
    plt.legend(bbox_to_anchor=(1, 0), loc="lower right", fontsize=20)
    plt.tight_layout()
    plt.savefig(RESULT_ROOT / "experiment1_energy_plot.png", dpi=300)
    plt.close()


def main():
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)
    CASE_ROOT.mkdir(parents=True, exist_ok=True)

    total_cases = len(TEST_LIST) * REPEAT_COUNT
    completed_cases = 0

    existing_completed_cases = 0
    for N, V, tot in TEST_LIST:
        for iterr in range(REPEAT_COUNT):
            payload_path = CASE_ROOT / case_name(N, V, tot, iterr) / "result.json"
            if not payload_path.exists():
                continue
            payload = json.loads(payload_path.read_text(encoding="utf-8"))
            if is_case_fully_completed(payload):
                existing_completed_cases += 1

    if existing_completed_cases == total_cases and total_cases > 0:
        print("[INFO] All experiment cases are already completed.")
        write_summary_files()
        print(f"[FINISH] Regenerated tables and figure under {RESULT_ROOT}")
        return

    for N, V, tot in TEST_LIST:
        for iterr in range(REPEAT_COUNT):
            payload = run_case(N, V, tot, iterr)
            if payload.get("status") == "completed":
                completed_cases += 1
            write_summary_files()
            print(f"[PROGRESS] {completed_cases}/{total_cases} cases completed")

    write_summary_files()
    print(f"[FINISH] Results saved under {RESULT_ROOT}")


if __name__ == "__main__":
    main()
